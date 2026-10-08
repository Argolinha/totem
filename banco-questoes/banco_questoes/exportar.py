"""Gravação do banco em CSV (';', UTF-8), XLSX e JSON com as 13 colunas."""
import csv
import json
from pathlib import Path

from .schema import COLUNAS


def gravar_csv(questoes: list, caminho: Path, bom: bool = False):
    with open(caminho, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUNAS, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        w.writerows(questoes)


def gravar_xlsx(questoes: list, caminho: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Questoes"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for q in questoes:
        linha = [q[c] for c in COLUNAS]
        linha[COLUNAS.index("Ano")] = int(q["Ano"]) if q["Ano"] else None
        ws.append(linha)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    larguras = {"Banca": 12, "Orgao": 14, "Ano": 7, "Disciplina": 30, "Assunto": 36, "Enunciado": 80, "Gabarito": 9}
    for i, nome in enumerate(COLUNAS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = larguras.get(nome, 40)
    wb.save(caminho)


def gravar_json(questoes: list, caminho: Path):
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump([{c: q[c] for c in COLUNAS} for q in questoes], f, ensure_ascii=False, indent=1)

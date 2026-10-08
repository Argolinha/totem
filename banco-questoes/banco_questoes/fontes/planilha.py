"""Importador genérico de CSV / XLSX / JSON já no formato das 13 colunas.

Serve para somar ao banco lotes vindos de outras fontes (provas transcritas,
bases licenciadas, questões autorais). Nomes de coluna são comparados sem
acento e sem diferenciar maiúsculas ("Órgão" == "Orgao"); alguns sinônimos
comuns também são aceitos.
"""
import csv
import json
import unicodedata
from pathlib import Path

from ..schema import COLUNAS

_SINONIMOS = {
    "orgao": "Orgao",
    "instituicao": "Orgao",
    "materia": "Disciplina",
    "topico": "Assunto",
    "pergunta": "Enunciado",
    "questao": "Enunciado",
    "resposta": "Gabarito",
    "comentario": "Comentario_Professor",
    "resolucao": "Comentario_Professor",
}
for _l in "abcde":
    _SINONIMOS[f"alternativa{_l}"] = f"Alternativa_{_l.upper()}"
    _SINONIMOS[f"opcao{_l}"] = f"Alternativa_{_l.upper()}"
    _SINONIMOS[_l] = f"Alternativa_{_l.upper()}"


def _chave(nome: str) -> str:
    t = unicodedata.normalize("NFKD", str(nome).strip().lower())
    return "".join(c for c in t if c.isalnum() and not unicodedata.combining(c))


_MAPA = {_chave(c): c for c in COLUNAS}
_MAPA.update({k: v for k, v in _SINONIMOS.items() if k not in _MAPA})


def _renomear(linha: dict) -> dict:
    saida = {}
    for k, v in linha.items():
        destino = _MAPA.get(_chave(k)) if k is not None else None
        if destino and not saida.get(destino):
            saida[destino] = "" if v is None else str(v)
    return saida


def ler(caminho: Path) -> list:
    caminho = Path(caminho)
    ext = caminho.suffix.lower()
    if ext == ".json":
        dados = json.loads(caminho.read_text(encoding="utf-8-sig"))
        return [_renomear(d) for d in dados]
    if ext == ".csv":
        texto = caminho.read_text(encoding="utf-8-sig")
        delim = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
        return [_renomear(d) for d in csv.DictReader(texto.splitlines(keepends=True), delimiter=delim)]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(caminho, read_only=True, data_only=True).active
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c) if c is not None else "" for c in next(linhas)]
        return [_renomear(dict(zip(cab, l))) for l in linhas if any(v not in (None, "") for v in l)]
    raise ValueError(f"formato não suportado: {caminho.name}")
